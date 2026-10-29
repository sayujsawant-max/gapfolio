"""Measure the Stage 1 evaluation targets and print one report.

    python scripts/evaluate.py

Reproduces every number quoted in the Stage 3 report that is not the threshold
sweep (that is scripts/tune_thresholds.py):

  - section extraction accuracy, on standard headings and on renamed ones
  - skill-entity extraction precision / recall / F1
  - ATS defect-injection detection rate and false-alarm rate
  - parse confidence on clean versus defective files
  - roadmap grounding: does the validator catch what an LLM could invent
  - end-to-end latency of the analysis pipeline

Everything is measured against single-annotator, synthetic fixtures, and the
report says so.
"""
from __future__ import annotations

import io
import json
import os
import pathlib
import statistics
import sys
import tempfile
import time
import warnings

warnings.filterwarnings("ignore")
ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["GAPFOLIO_DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp()}/eval.db"
os.environ.setdefault("GAPFOLIO_EMBEDDING_PROVIDER", "onnx")
os.environ.setdefault("GAPFOLIO_EMBEDDING_DIM", "384")
os.environ["GAPFOLIO_ANTHROPIC_API_KEY"] = ""

import docx  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from fpdf import FPDF  # noqa: E402

from app.db import SessionLocal, init_db  # noqa: E402
from app.services import entities, matching, vectorstore  # noqa: E402
from app.services.embeddings import effective_provider  # noqa: E402
from app.services.parsing import ParseError, ats_checks, parse_document, parse_text  # noqa: E402

FIX = ROOT / "tests" / "fixtures"
PAIRS = json.loads((FIX / "labeled_pairs.json").read_text(encoding="utf-8"))
ANN = json.loads((FIX / "annotations.json").read_text(encoding="utf-8"))
RESUMES: dict[str, str] = PAIRS["resumes"]


def header(title: str) -> None:
    print(f"\n=== {title} ===")


# ------------------------------------------------------------------ documents

def wrap(key: str, variant: str | None):
    """Full resume text with a header block and an education section, headings
    optionally renamed. Returns (text, gold body lines, gold headings)."""
    rename = ANN["variants"].get(variant, {}) if variant else {}
    first = key.split("_")[0].title()
    out: list[str] = []
    body: list[tuple[str, str]] = []
    heads: list[tuple[str, str]] = []

    for line in (f"{first.upper()} RAO", f"{first.lower()}@example.edu | +91 98765 43210 | github.com/{first.lower()}"):
        out.append(line)
        body.append((line, "header"))

    def add_heading(std: str) -> None:
        shown = rename.get(std, std)
        out.append(shown)
        heads.append((shown, ANN["headings"][std]))

    add_heading("EDUCATION")
    edu = "B.Tech Computer Science, 2027, CGPA 8.1 at Pune Institute of Technology"
    out.append(edu)
    body.append((edu, "education"))

    section = None
    for raw in RESUMES[key].splitlines():
        line = raw.strip()
        if not line:
            continue
        if line in ANN["headings"]:
            section = ANN["headings"][line]
            add_heading(line)
        else:
            out.append(line)
            body.append((line, section))
    return "\n".join(out), body, heads


# ------------------------------------------------------------------ 1. sections

def eval_sections() -> None:
    header("1. Section extraction")
    print(f"embedding provider: {effective_provider()}")
    results = {}
    for variant in (None, "rulebook", "creative_a", "creative_b"):
        for fallback in (True, False):
            lines_ok = lines_n = heads_ok = heads_n = 0
            for key in RESUMES:
                text, body, heads = wrap(key, variant)
                segs = entities.segment_sections(text, fallback=fallback)
                predicted = {ln: s.section for s in segs for ln in s.lines}
                for ln, gold in body:
                    lines_n += 1
                    lines_ok += predicted.get(ln) == gold
                for h, gold in heads:
                    heads_n += 1
                    heads_ok += any(s.heading == h and s.section == gold for s in segs)
            results[(variant or "standard", fallback)] = (lines_ok, lines_n, heads_ok, heads_n)
    print(f"{'headings':<12} {'fallback':<9} {'line accuracy':<20} heading accuracy")
    for (variant, fb), (lo, ln, ho, hn) in results.items():
        print(f"{variant:<12} {'on' if fb else 'off':<9} {lo:>3}/{ln:<4} = {lo/ln:6.1%}      {ho:>3}/{hn:<3} = {ho/hn:6.1%}")
    n_docs = len(RESUMES) * 4
    tot = [sum(v[i] for (var, fb), v in results.items() if fb) for i in range(4)]
    print(f"overall with fallback on: lines {tot[0]}/{tot[1]} = {tot[0]/tot[1]:.1%} across {n_docs} documents")


# ------------------------------------------------------------------ 2. skills

def eval_skills() -> None:
    header("2. Skill-entity extraction")
    tp = fp = fn = 0
    vocab = set(entities.SKILLS)
    in_vocab_gold = in_vocab_hit = 0
    misses, extras = [], []
    for key, text in RESUMES.items():
        gold = set(ANN["skills"][key])
        pred = set(entities.extract_skills(text))
        tp += len(gold & pred)
        fp += len(pred - gold)
        fn += len(gold - pred)
        in_vocab_gold += len(gold & vocab)
        in_vocab_hit += len(gold & vocab & pred)
        misses += [f"{key}:{s}" for s in sorted(gold - pred)]
        extras += [f"{key}:{s}" for s in sorted(pred - gold)]
    p, r = tp / (tp + fp), tp / (tp + fn)
    f1 = 2 * p * r / (p + r)
    total_gold = tp + fn
    print(f"gold skills {total_gold} (of which {total_gold - in_vocab_gold} are outside the vocabulary)")
    print(f"precision {p:.3f}  recall {r:.3f}  F1 {f1:.3f}")
    print(f"recall on in-vocabulary skills only: {in_vocab_hit}/{in_vocab_gold} = {in_vocab_hit/in_vocab_gold:.3f}")
    print("missed (out-of-vocabulary):", ", ".join(misses))
    print("false positives:", ", ".join(extras) or "none")


# ------------------------------------------------------------------ 3. ATS defects

BODY = [
    "Built a library management system in Django used by 300 students",
    "Reduced page load time by 35% with query caching",
    "Automated weekly reports in Python saving 6 hours per week",
    "Wrote 25 unit tests raising coverage to 80% on the billing module",
    "Led a team of 4 to ship a campus events app to 900 users",
    "Migrated a MySQL database of 50k rows with zero downtime",
]


def lines_for(key: str) -> list[str]:
    first = key.split("_")[0].title()
    base = [l.strip() for l in RESUMES[key].splitlines() if l.strip()]
    return [f"{first.upper()} RAO  {first.lower()}@example.edu  +91 98765 43210", "EDUCATION",
            "B.Tech Computer Science, 2027, CGPA 8.1"] + base + ["- " + b for b in BODY]


def pdf_single(lines):
    p = FPDF(); p.add_page(); p.set_font("helvetica", size=10)
    for l in lines:
        p.set_x(p.l_margin); p.multi_cell(w=p.epw, h=6, text=l)
    return bytes(p.output())


def pdf_two_col(lines):
    p = FPDF(); p.add_page(); p.set_font("helvetica", size=10)
    half = (len(lines) + 1) // 2
    for i, l in enumerate(lines[:half]):
        p.set_xy(12, 20 + i * 7); p.cell(90, 6, l[:48])
    for i, l in enumerate(lines[half:]):
        p.set_xy(110, 20 + i * 7); p.cell(90, 6, l[:48])
    return bytes(p.output())


def pdf_scan():
    p = FPDF(); p.add_page()
    return bytes(p.output())


def pdf_unembedded(lines):
    """A hand-built PDF that names Calibri but carries no font file."""
    esc = lambda s: s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = "BT /F1 10 Tf 40 800 Td 12 TL " + " ".join(f"({esc(l[:90])}) Tj T*" for l in lines) + " ET"
    widths = " ".join(["500"] * 95)
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        f"<< /Type /Font /Subtype /TrueType /BaseFont /Calibri /FirstChar 32 /LastChar 126 /Widths [{widths}] "
        "/FontDescriptor 6 0 R /Encoding /WinAnsiEncoding >>",
        "<< /Type /FontDescriptor /FontName /Calibri /Flags 32 /FontBBox [0 -200 1000 900] "
        "/ItalicAngle 0 /Ascent 900 /Descent -200 /CapHeight 700 /StemV 80 >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def docx_bytes(lines, table=False, header_contact=False, textbox=False):
    d = docx.Document()
    body = list(lines)
    if header_contact:
        d.sections[0].header.paragraphs[0].text = body.pop(0)
    if table:
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "SKILLS"; t.cell(0, 1).text = "Python, SQL, Git"
        t.cell(1, 0).text = "EDUCATION"; t.cell(1, 1).text = "B.Tech 2027"
    for l in body:
        d.add_paragraph(l)
    if textbox:
        ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
              'xmlns:v="urn:schemas-microsoft-com:vml"')
        run = parse_xml(
            f'<w:r {ns}><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Skills: Python, SQL</w:t></w:r></w:p>'
            "</w:txbxContent></v:textbox></v:shape></w:pict></w:r>")
        d.add_paragraph()._p.append(run)
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


def flagged(checks, part) -> bool:
    return any(part in c["check"] and c["status"] in ("fail", "warn") for c in checks)


def eval_ats() -> None:
    header("3. ATS audit: defect injection")
    caught: dict[str, list[bool]] = {}
    clean_alarms = []
    confidences: dict[str, list[float]] = {"clean": [], "two-column": [], "table": [], "text box": [], "header contact": []}

    def record(defect, ok):
        caught.setdefault(defect, []).append(bool(ok))

    for key in RESUMES:
        lines = lines_for(key)

        r = parse_document(pdf_single(lines), "clean.pdf")
        checks = ats_checks(r.meta)
        clean_alarms.append([c["check"] for c in checks if c["group"] == "parse" and c["status"] != "pass"])
        confidences["clean"].append(r.meta["parse_confidence"])

        r = parse_document(pdf_two_col(lines), "two.pdf")
        record("Multi-column layout", flagged(ats_checks(r.meta), "Multi-column"))
        confidences["two-column"].append(r.meta["parse_confidence"])

        try:
            parse_document(pdf_scan(), "scan.pdf"); record("Scanned PDF (no text layer)", False)
        except ParseError:
            record("Scanned PDF (no text layer)", True)

        r = parse_document(docx_bytes(lines, table=True), "t.docx")
        record("Layout table", flagged(ats_checks(r.meta), "Tables"))
        confidences["table"].append(r.meta["parse_confidence"])

        r = parse_document(docx_bytes(lines, textbox=True), "tb.docx")
        record("Text box", flagged(ats_checks(r.meta), "Text boxes"))
        confidences["text box"].append(r.meta["parse_confidence"])

        r = parse_document(docx_bytes(lines, header_contact=True), "h.docx")
        record("Contact in page header", flagged(ats_checks(r.meta), "header"))
        confidences["header contact"].append(r.meta["parse_confidence"])

        r = parse_text("\n".join(l.split("  ")[0] if "@" in l else l for l in lines))
        record("Missing email in body", flagged(ats_checks(r.meta), "No email"))

        renamed = [ANN["variants"]["creative_a"].get(l, l) for l in lines]
        r = parse_text("\n".join(renamed))
        record("Non-standard section headings", flagged(ats_checks(r.meta), "Non-standard"))

        # several defects at once: two columns, no contact details, odd headings
        stripped = [ANN["variants"]["creative_a"].get(l, l) for l in lines[1:]]
        r = parse_document(pdf_two_col(stripped), "bad.pdf")
        confidences.setdefault("compound defects", []).append(r.meta["parse_confidence"])

        r = parse_document(pdf_unembedded(lines), "font.pdf")
        record("Font not embedded", flagged(ats_checks(r.meta), "Fonts not embedded"))

    n = len(RESUMES)
    total = hit = 0
    for defect, oks in caught.items():
        print(f"  {defect:<34} {sum(oks)}/{len(oks)}")
        total += len(oks); hit += sum(oks)
    print(f"  OVERALL {hit}/{total} = {hit/total:.1%}")
    alarms = [a for a in clean_alarms if a]
    print(f"  false alarms on {n} clean single-column PDFs: {len(alarms)}/{n}  {alarms[:2]}")
    print("  mean parse confidence:", {k: round(statistics.mean(v), 2) for k, v in confidences.items()})
    print("  parse confidence flagged 'low' (<0.60) on:",
          {k: sum(c < 0.6 for c in v) for k, v in confidences.items()})


# ------------------------------------------------------------------ 4. grounding

def eval_grounding() -> None:
    header("4. Roadmap grounding")
    init_db(); db = SessionLocal()
    from app.services.parsing import parse_text as pt

    caught = {"invented gap": [0, 0], "invented number": [0, 0], "unflagged bullet": [0, 0]}
    kept_legit = [0, 0]
    ordered = [0, 0]
    for key, text in RESUMES.items():
        rid = f"g-{key}"
        vectorstore.index_document(db, "resume", rid, text)
        reqs = [{"label": c["label"], "jd_quote": c["quote"], "weight": "core"} for c in PAIRS["cases"] if c["resume"] == key]
        ms = matching.match_requirements(db, rid, reqs)
        ranked = matching.rank_gaps(ms)[:5]
        bullets = [matching._strip_bullet(b) for b in pt(text + "\n- Helped the team with testing tasks\n- Worked on the backend").meta["weak_verb_bullets"]][:6]
        fb = matching._fallback_roadmap(ranked, bullets)
        if not ranked or not bullets:
            continue
        legit_gap, legit_before = ranked[0]["label"], bullets[0]
        data = {
            "roadmap": [
                {"skill": legit_gap, "action": "legit", "proof": "Built [N]"},
                {"skill": "Quantum Computing", "action": "invented", "proof": "x"},
            ],
            "rewrites": [
                {"before": legit_before, "after": "Delivered the feature, measured at [N]", "why": "ok"},
                {"before": legit_before, "after": "Delivered the feature for 5000 users", "why": "invented number"},
                {"before": "Single-handedly built the entire platform", "after": "Built it", "why": "not in the resume"},
            ],
            "questions": ["q1", "q2", "q3"],
        }
        out = matching._ground(data, ranked, bullets, fb)
        skills = [r["skill"] for r in out["roadmap"]]
        befores = [rw["after"] for rw in out["rewrites"]]
        caught["invented gap"][1] += 1; caught["invented gap"][0] += "Quantum Computing" not in skills
        caught["invented number"][1] += 1; caught["invented number"][0] += not any("5000" in a for a in befores)
        caught["unflagged bullet"][1] += 1
        caught["unflagged bullet"][0] += not any(rw["before"].startswith("Single-handedly") for rw in out["rewrites"])
        kept_legit[1] += 1; kept_legit[0] += legit_gap in skills and any("[N]" in a for a in befores)
        scores = [r["priority_score"] for r in out["roadmap"]]
        ordered[1] += 1; ordered[0] += scores == sorted(scores, reverse=True)
    for k, (a, b) in caught.items():
        print(f"  invented content caught - {k:<18} {a}/{b}")
    print(f"  legitimate items kept: {kept_legit[0]}/{kept_legit[1]}")
    print(f"  roadmap in descending priority order: {ordered[0]}/{ordered[1]}")

    # the prompt must not contain resume text
    captured = {}
    real = matching.complete_json
    matching.complete_json = lambda prompt, max_tokens=0: captured.setdefault("p", prompt) and {}
    text = RESUMES["backend_fresher"]
    vectorstore.index_document(db, "resume", "g-prompt", text)
    reqs = [{"label": c["label"], "jd_quote": c["quote"], "weight": "core"} for c in PAIRS["cases"] if c["resume"] == "backend_fresher"]
    ms = matching.match_requirements(db, "g-prompt", reqs)
    matching.build_roadmap("Backend Engineer", ms, ["Helped the team with testing tasks"], {})
    matching.complete_json = real
    prompt = captured.get("p", "")
    body_lines = [l[2:] for l in text.splitlines() if l.startswith("- ")]
    leaked = [l for l in body_lines if l in prompt]
    print(f"  prompt sent to the LLM: {len(prompt)} chars; full resume bullets appearing in it: "
          f"{len(leaked)} of {len(body_lines)} (the resume itself is never included)")


# ------------------------------------------------------------------ 5. latency

def eval_latency() -> None:
    header("5. Latency (parse + index + match + score + ATS + roadmap, warm model)")
    init_db(); db = SessionLocal()
    t0 = time.perf_counter()
    from app.services.embeddings import embed_one
    embed_one("warm up")
    print(f"  model load / first embedding: {time.perf_counter() - t0:.1f} s")
    times = []
    for key, text in RESUMES.items():
        reqs = [{"label": c["label"], "jd_quote": c["quote"], "weight": "core"} for c in PAIRS["cases"] if c["resume"] == key]
        for rep in range(4):
            t = time.perf_counter()
            res = parse_document(pdf_single(lines_for(key)), "r.pdf")
            rid = f"lat-{key}-{rep}"
            vectorstore.index_document(db, "resume", rid, res.text)
            ms = matching.match_requirements(db, rid, reqs)
            matching.score(ms, {})
            ats_checks(res.meta)
            matching.build_roadmap("Role", ms, res.meta.get("weak_verb_bullets"), {})
            times.append(time.perf_counter() - t)
    times.sort()
    print(f"  {len(times)} runs: median {statistics.median(times)*1000:.0f} ms, "
          f"p95 {times[int(0.95 * len(times)) - 1]*1000:.0f} ms, max {times[-1]*1000:.0f} ms")


if __name__ == "__main__":
    print(f"provider in use: {effective_provider()}")
    for fn in (eval_sections, eval_skills, eval_ats, eval_grounding, eval_latency):
        fn()
