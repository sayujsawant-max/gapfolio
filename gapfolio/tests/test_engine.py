"""Engine-level tests: parsing signals, match semantics, calibration maths."""
import io
import os
import tempfile

import pytest

os.environ.setdefault("GAPFOLIO_DATABASE_URL", f"sqlite:///{tempfile.mkdtemp()}/engine.db")
os.environ["GAPFOLIO_EMBEDDING_PROVIDER"] = "hash"
os.environ["GAPFOLIO_EMBEDDING_DIM"] = "1024"  # the offline stub was tuned at 1024-d

from app.db import SessionLocal, init_db  # noqa: E402
from app.models import Outcome, SkillWeight  # noqa: E402
from app.services import calibration, matching, vectorstore  # noqa: E402
from app.services.parsing import ParseError, ats_checks, parse_document, parse_text  # noqa: E402


# --- parsing --------------------------------------------------------------

def _pdf_bytes(lines, two_column=False):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=11)
    if two_column:
        for i, line in enumerate(lines):
            pdf.set_xy(15 if i % 2 == 0 else 110, 20 + (i // 2) * 8)
            pdf.cell(85, 6, line)
    else:
        for line in lines:
            pdf.set_x(pdf.l_margin); pdf.multi_cell(w=pdf.epw, h=6, text=line)
    return bytes(pdf.output())


RESUME_LINES = [
    "PRIYA DESHMUKH  priya@example.edu  +91 98765 43210",
    "EDUCATION",
    "B.Tech Electronics, 2027, CGPA 8.1",
    "PROJECTS",
    "- Built a REST API in Python handling 2000 requests per minute",
    "- Helped the team with testing tasks",
    "EXPERIENCE",
    "- Fixed 12 defects found during code review",
    "SKILLS",
    "Python, Docker, PostgreSQL, Git, AWS",
] * 3


def test_pdf_text_extraction():
    result = parse_document(_pdf_bytes(RESUME_LINES), "resume.pdf")
    assert result.source == "pdf"
    assert "PRIYA DESHMUKH" in result.text
    assert result.meta["has_email"] is True
    assert result.meta["has_phone"] is True
    assert result.meta["is_scanned"] is False


def test_scanned_pdf_is_rejected_with_a_useful_message():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()  # no text at all — stands in for a scan
    with pytest.raises(ParseError) as e:
        parse_document(bytes(pdf.output()), "scan.pdf")
    assert "no text layer" in str(e.value)


def test_docx_table_layout_is_flagged():
    import docx

    doc = docx.Document()
    doc.add_paragraph("ROHIT SHARMA  rohit@example.edu")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "EDUCATION"
    table.cell(0, 1).text = "B.Tech CSE 2027"
    table.cell(1, 0).text = "SKILLS"
    table.cell(1, 1).text = "Python, Docker, Kubernetes, Git, SQL"
    for _ in range(6):
        doc.add_paragraph("- Built a service that cut latency by 40% for 5000 users")
    buf = io.BytesIO()
    doc.save(buf)

    result = parse_document(buf.getvalue(), "resume.docx")
    assert result.meta["layout_tables"] is True
    checks = ats_checks(result.meta)
    assert any(c["check"] == "Tables in the layout" and c["status"] == "fail" for c in checks)


def test_weak_verbs_and_quantification_are_measured():
    result = parse_text(
        "PROJECTS\n"
        "- Helped the team with website tasks\n"
        "- Worked on the backend\n"
        "- Shipped an API serving 1200 users with 99.9% uptime\n"
        "EDUCATION\nB.Tech CSE\nSKILLS\nPython, Git\n"
    )
    assert len(result.meta["weak_verb_bullets"]) == 2
    assert result.meta["quantified_bullets"] == 1
    checks = {c["check"]: c for c in ats_checks(result.meta)}
    assert any("Weak opening verbs" in k for k in checks)


# --- matching -------------------------------------------------------------

@pytest.fixture(scope="module")
def db():
    init_db()
    s = SessionLocal()
    yield s
    s.close()


def test_listed_skill_is_partial_not_proven(db):
    """The rule the whole score rests on: a skill in a comma list is a claim,
    a skill inside an achievement is proof."""
    resume = (
        "PROJECTS\n"
        "- Containerised the payments service with Docker and shipped it via GitHub Actions\n"
        "EDUCATION\nB.Tech CSE 2027\n"
        "SKILLS\nKubernetes, Terraform, Python\n"
    )
    vectorstore.index_document(db, "resume", "r-test-1", resume)

    reqs = [
        {"label": "Docker containerisation", "jd_quote": "Containerise services with Docker", "weight": "core"},
        {"label": "Kubernetes", "jd_quote": "Deploy workloads on Kubernetes", "weight": "core"},
        {"label": "Salesforce administration", "jd_quote": "Administer Salesforce CRM", "weight": "core"},
    ]
    by_label = {m.label: m for m in matching.match_requirements(db, "r-test-1", reqs)}

    assert by_label["Docker containerisation"].status == "proven"
    assert by_label["Kubernetes"].status == "partial"
    assert by_label["Kubernetes"].source_kind == "skill_listed"
    assert by_label["Salesforce administration"].status == "missing"


def test_score_ordering_and_readiness(db):
    def m(status, weight="core"):
        return matching.Matched("x", "", weight, status, None, 0.5, "")

    strong, _ = matching.score([m("proven")] * 8, {})
    mixed, _ = matching.score([m("proven")] * 4 + [m("partial")] * 4, {})
    weak, readiness = matching.score([m("missing")] * 8, {})

    assert strong == 100.0
    assert weak == 0.0
    assert strong > mixed > weak
    assert readiness == "Not yet"


def test_learned_weights_change_the_score(db):
    def m(label, status):
        return matching.Matched(label, "", "core", status, None, 0.5, "")

    matches = [m("docker", "missing"), m("figma", "proven")]
    neutral, _ = matching.score(matches, {})
    # Docker predicts callbacks strongly, so missing it should hurt more.
    weighted, _ = matching.score(matches, {"docker": 1.6, "figma": 0.5})
    assert weighted < neutral


# --- calibration ----------------------------------------------------------

def test_calibration_learns_lift_from_outcomes(db):
    db.query(Outcome).delete()
    db.query(SkillWeight).delete()
    db.commit()

    def snap(docker_status):
        return [
            {"label": "docker", "status": docker_status, "weight": "core"},
            {"label": "figma", "status": "proven", "weight": "core"},
        ]

    # Twelve applicants with Docker proven mostly get callbacks; twelve without
    # mostly don't. Figma is proven throughout, so it should carry no lift.
    for i in range(12):
        db.add(Outcome(report_id=f"p{i}", status="interview" if i < 10 else "rejected",
                       score_at_apply=70, requirement_snapshot=snap("proven")))
    for i in range(12):
        db.add(Outcome(report_id=f"n{i}", status="interview" if i < 2 else "rejected",
                       score_at_apply=40, requirement_snapshot=snap("missing")))
    db.commit()

    result = calibration.recalibrate(db, "backend")
    assert result["skills_updated"] >= 1

    docker = db.query(SkillWeight).filter_by(skill="docker", role_family="backend").one()
    assert docker.lift > 0.4, "Docker should show a large positive lift"
    assert docker.weight > 1.0, "and should therefore be weighted up"
    # Figma never varies, so it has no comparison group and stays unweighted.
    assert db.query(SkillWeight).filter_by(skill="figma").first() is None

    ev = calibration.evaluate(db)
    assert ev["n"] == 24
    assert ev["point_biserial_r"] > 0.25
    assert "tracks callbacks" in ev["verdict"]


def test_evaluation_refuses_to_claim_predictiveness_without_data(db):
    db.query(Outcome).delete()
    db.commit()
    assert "Not enough outcomes" in calibration.evaluate(db)["verdict"]


# --- calibrated accuracy --------------------------------------------------

def test_matcher_agrees_with_human_labels(db):
    """Regression guard on the labeled set. If a change to chunking, the
    ontology or the thresholds drops agreement below the bar, this fails
    before anyone sees a worse report."""
    import json
    import pathlib

    from app.services import lexical, vectorstore
    from app.services.embeddings import cosine, embed_one
    from app.services.matching import _resume_chunks, DENSE_W, LEXICAL_W, PROVEN_T, PARTIAL_T

    data = json.loads(
        (pathlib.Path(__file__).parent / "fixtures" / "labeled_pairs.json").read_text()
    )
    for key, text in data["resumes"].items():
        vectorstore.index_document(db, "resume", f"fx-{key}", text)

    correct = false_missing = 0
    for case in data["cases"]:
        chunks = _resume_chunks(db, f"fx-{case['resume']}")
        idf = lexical.build_idf([c.chunk for c in chunks])
        query = f"{case['label']}. {case['quote']}"
        qvec = embed_one(query)

        best = proof = 0.0
        for c in chunks:
            rel = DENSE_W * max(cosine(qvec, c.vector), 0.0) + LEXICAL_W * lexical.coverage(
                query, c.chunk, idf
            )
            best = max(best, rel)
            if c.kind in ("achievement", "context"):
                proof = max(proof, rel)

        got = "proven" if proof >= PROVEN_T else "partial" if best >= PARTIAL_T else "missing"
        correct += got == case["expect"]
        if got == "missing" and case["expect"] != "missing":
            false_missing += 1

    n = len(data["cases"])
    assert correct / n >= 0.80, f"agreement fell to {correct / n:.0%} over {n} cases"
    # Telling a student to learn something they already have is the error that
    # destroys trust fastest, so it gets a tighter bound than overall accuracy.
    assert false_missing / n <= 0.10, f"false 'missing' rate {false_missing / n:.0%} too high"
